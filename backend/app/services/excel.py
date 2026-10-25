"""Локальное Excel-зеркало данных: data/mirror.xlsx (Orders / Stock / Operations).

Обновляется после каждой операции с заявкой и доступно на скачивание
через /api/export/excel. Когда подключены Google Sheets — то же самое
падает в таблицу; Excel-зеркало работает и без внешних ключей.
"""
import logging
from pathlib import Path

from ..database import SessionLocal
from ..models import Operation, Order, Product

DATA_DIR = Path(__file__).resolve().parents[3] / "data"
MIRROR = DATA_DIR / "mirror.xlsx"

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ORDERS_HEADERS = ("Номер", "Дата", "Клиент", "Контакт", "Товар", "Кол-во",
                  "Цена", "Сумма", "Статус", "Канал", "Создал")
STOCK_HEADERS = ("SKU", "Товар", "Остаток", "Цена", "Обновлён")
OPS_HEADERS = ("Время", "Актёр", "Действие", "Сущность", "ID", "Детали", "Причина")


def mirror_path() -> Path:
    return MIRROR


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    return ws


def sync_mirror(db) -> str:
    """Пишет актуальное зеркало. Возвращает путь к файлу."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1D1D1F")

    def style(ws):
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(42, max(9, width + 2))

    orders = db.query(Order).order_by(Order.created_at.desc()).limit(1000).all()
    ws_o = wb.create_sheet("Orders")
    ws_o.append(ORDERS_HEADERS)
    for o in orders:
        ws_o.append([o.number, o.created_at.strftime("%Y-%m-%d %H:%M"),
                     o.customer_name, o.contact, o.product.name, o.qty,
                     o.price, o.amount, o.status, o.source,
                     o.creator.full_name or o.creator.login if o.creator else "—"])
    style(ws_o)

    ws_s = wb.create_sheet("Stock")
    ws_s.append(STOCK_HEADERS)
    for p in db.query(Product).order_by(Product.sku).all():
        ws_s.append([p.sku, p.name, p.stock, p.price, p.updated_at.strftime("%Y-%m-%d %H:%M")])
    style(ws_s)

    ops = db.query(Operation).order_by(Operation.ts.desc()).limit(500).all()
    ws_a = wb.create_sheet("Operations")
    ws_a.append(OPS_HEADERS)
    for op in ops:
        ws_a.append([op.ts.strftime("%Y-%m-%d %H:%M"), op.actor, op.action,
                     op.entity, op.entity_id, op.details, op.reason])
    style(ws_a)

    DATA_DIR.mkdir(exist_ok=True)
    wb.save(MIRROR)
    return str(MIRROR)


def sync_in_background():
    """Фоновая задача: зеркалит БД в Excel собственным соединением."""
    db = SessionLocal()
    try:
        sync_mirror(db)
    except Exception as e:  # Excel не должен ронять бизнес-операцию
        logging.warning("Excel mirror failed: %s", e)
    finally:
        db.close()
